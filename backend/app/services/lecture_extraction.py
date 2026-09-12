"""
Lecture file extraction service — Phase 7, Sub-feature 7.1.

Provides:
  - extract_pptx_content(file_path) -> dict   never raises
  - validate_lecture_filename(filename) -> str | None   never raises

Exception-safe by design, matching notebook.py's convention: any internal
failure is caught and communicated through the return value's "valid"/"error"
keys rather than propagating to callers.
"""

import logging
from pathlib import Path
from typing import Union

logger = logging.getLogger(__name__)

LEGACY_PPT_ERROR = (
    "Legacy .ppt format is not supported — please save as .pptx and re-upload."
)


def validate_lecture_filename(filename: str) -> str | None:
    """
    Check the filename's extension before any extraction is attempted.
    Returns an error message string if the file should be rejected, or None
    if it's an acceptable .pptx.

    Does not rely on the client-supplied Content-Type header — matching the
    rest of this project's uploads, which gate on filename/extension, not on
    an unvalidated client header.
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".pptx":
        return None
    if suffix == ".ppt":
        return LEGACY_PPT_ERROR
    return (
        f"Unsupported file type '{suffix or '(none)'}' — only .pptx lecture "
        f"files are accepted."
    )


def extract_pptx_content(file_path: Union[str, Path]) -> dict:
    """
    Parse a .pptx file and return a structured per-slide result.

    Returns
    -------
    {
        "slides": list[dict],   # [{"slide_number": int, "slide_text": str,
                                  #   "notes_text": str}, ...] in slide order.
                                  # slide_text includes any table-derived text,
                                  # prefixed with "[Table]" per table. notes_text
                                  # is "" when the slide has no notes (never an
                                  # error).
        "valid": bool,
        "error": str | None,
    }

    This function NEVER raises. Any failure (file not found, corrupt/non-pptx
    binary, etc.) is captured and returned as valid=False with a descriptive
    error string, matching notebook.py's established convention.
    """
    result: dict = {"slides": [], "valid": False, "error": None}

    try:
        path = Path(file_path)

        if not path.exists():
            result["error"] = f"File not found: {file_path}"
            return result

        if path.stat().st_size == 0:
            result["error"] = f"File is empty: {file_path}"
            return result

        from pptx import Presentation
        from pptx.exc import PackageNotFoundError

        try:
            prs = Presentation(str(path))
        except PackageNotFoundError:
            result["error"] = (
                f"File does not appear to be a valid .pptx package (corrupted "
                f"or not actually PowerPoint format): {file_path}"
            )
            return result

        slides: list[dict] = []
        for slide_number, slide in enumerate(prs.slides, start=1):
            slide_text = _extract_slide_text(slide)
            notes_text = _extract_notes_text(slide)
            slides.append(
                {
                    "slide_number": slide_number,
                    "slide_text": slide_text,
                    "notes_text": notes_text,
                }
            )

        result["slides"] = slides
        result["valid"] = True
        return result

    except Exception as exc:  # noqa: BLE001 — deliberate catch-all, see docstring
        logger.warning("Failed to extract .pptx content from %s: %s", file_path, exc)
        result["error"] = f"Failed to parse .pptx file: {exc}"
        return result


# ── Internal helpers ─────────────────────────────────────────────────────────

def _extract_slide_text(slide) -> str:
    """
    Extract all text from all shapes on a slide, in shape order.

    Table shapes require explicit separate handling: a shape containing a
    table does not expose its content via shape.text_frame, and a table
    shape / a text-frame shape are mutually exclusive in python-pptx — so
    shape.has_table is checked for EVERY shape, not as a fallback only when
    text_frame extraction is empty. Non-table, non-text shapes (e.g. images
    with no text) are skipped cleanly, not treated as errors.
    """
    parts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            parts.append(_extract_table_text(shape.table))
        elif getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text
            if text and text.strip():
                parts.append(text.strip())
    return "\n".join(p for p in parts if p)


def _extract_table_text(table) -> str:
    """
    Serialize a table's cell text: cells within a row joined with " | ",
    rows separated by newlines, prefixed with a "[Table]" marker so
    table-derived text is clearly labeled within the slide's combined text.
    """
    rows_text = []
    for row in table.rows:
        cells_text = [cell.text.strip() for cell in row.cells]
        rows_text.append(" | ".join(cells_text))
    return "[Table]\n" + "\n".join(rows_text)


def _extract_notes_text(slide) -> str:
    """Extract speaker notes text, or "" if the slide has no notes slide/text."""
    if not slide.has_notes_slide:
        return ""
    notes_slide = slide.notes_slide
    if notes_slide is None or notes_slide.notes_text_frame is None:
        return ""
    return notes_slide.notes_text_frame.text.strip()


# ── Chunking ──────────────────────────────────────────────────────────────────

# Deliberate, stated choice (not an LLM prompt, so no sign-off needed): slide
# content is typically short — a handful of bullet points plus a paragraph of
# notes — so 800 chars (~130-160 words) keeps a typical slide as ONE chunk,
# while still splitting unusually dense notes into overlapping windows for
# embedding-friendly granularity in 7.2. 150 chars (~19%) of overlap preserves
# context across a split without excessive duplication.
CHUNK_SIZE = 800
CHUNK_OVERLAP = 150


def chunk_slides(slides: list[dict]) -> list[dict]:
    """
    Turn extract_pptx_content()'s per-slide result into chunk records.

    Slide text and notes text are kept as SEPARATE chunk-eligible pools (not
    combined into one blob per slide) — a chunk is scoped to exactly one
    (slide_number, source) pair and never crosses a slide boundary, which is
    a stricter guarantee than "never crosses a slide boundary" alone. This
    gives 7.2's future citations unambiguous granularity, e.g. "Slide 14,
    speaker notes" vs. "Slide 14, slide text", never a length-cut that mixes
    the two or spans two slides.

    Returns a flat list of:
        {"slide_number": int, "source": "slide_text" | "notes",
         "chunk_index": int, "chunk_text": str}
    Empty text (e.g. a slide with no notes) produces zero chunks for that pool.
    """
    chunks: list[dict] = []
    for slide in slides:
        slide_number = slide["slide_number"]
        for source, text in (
            ("slide_text", slide["slide_text"]),
            ("notes", slide["notes_text"]),
        ):
            for chunk_index, chunk_text in enumerate(_split_text(text)):
                chunks.append(
                    {
                        "slide_number": slide_number,
                        "source": source,
                        "chunk_index": chunk_index,
                        "chunk_text": chunk_text,
                    }
                )
    return chunks


def _split_text(text: str) -> list[str]:
    """
    Split text into CHUNK_SIZE-character windows with CHUNK_OVERLAP overlap.
    Text shorter than one chunk produces exactly one chunk. Empty/whitespace
    text produces zero chunks.
    """
    text = text.strip()
    if not text:
        return []
    if len(text) <= CHUNK_SIZE:
        return [text]

    pieces: list[str] = []
    start = 0
    step = CHUNK_SIZE - CHUNK_OVERLAP
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        pieces.append(text[start:end])
        if end == len(text):
            break
        start += step
    return pieces
