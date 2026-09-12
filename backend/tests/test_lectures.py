"""
Tests for POST/GET /api/v1/sessions/{session_id}/lectures and
GET .../lectures/{lecture_id}/download — Phase 7, Sub-feature 7.1.

Mirrors test_assignments.py's fixture/helper conventions (in-memory SQLite
with FK enforcement, self-contained per-file fixtures, storage functions
mocked to avoid real disk I/O in unit tests).

Run with:
    cd backend
    python -m pytest tests/test_lectures.py -v
"""

import io
from unittest.mock import patch

import pytest
from pptx import Presentation
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models.lecture_file import LectureFile
from app.models.lecture_chunk import LectureChunk
from app.models.session import LMSSession
from app.models.user import User, UserRole
from app.services.auth import create_access_token


# ===========================================================================
# Fixture builders
# ===========================================================================

def _make_pptx_bytes(with_notes: bool = True) -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Week 1 Intro"
    s.placeholders[1].text = "Welcome to the course"
    if with_notes:
        s.notes_slide.notes_text_frame.text = "Say hello and set expectations."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def _mock_save_lecture(session_id: int, filename: str, data: bytes) -> str:
    return f"{session_id}/lectures/{filename}"


def _post(client, token, session_id, filename, data, content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"):
    return client.post(
        f"/api/v1/sessions/{session_id}/lectures",
        files={"file": (filename, data, content_type)},
        headers={"Authorization": f"Bearer {token}"},
    )


def _post_mocked(client, token, session_id, filename, data, **kw):
    with patch("app.routers.lectures.save_lecture_file", side_effect=_mock_save_lecture):
        return _post(client, token, session_id, filename, data, **kw)


# ===========================================================================
# Fixtures
# ===========================================================================

@pytest.fixture()
def db():
    import app.models  # noqa: F401

    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _enable_foreign_keys(dbapi_conn, _connection_record):
        cursor = dbapi_conn.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine)
    TestingSession = sessionmaker(bind=engine)
    session = TestingSession()
    yield session
    session.close()
    Base.metadata.drop_all(engine)


@pytest.fixture()
def client(db):
    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db

    instructor = User(
        id=1, name="Prof Test", email="prof@test.com",
        hashed_password="hash", role=UserRole.instructor,
    )
    instructor2 = User(
        id=2, name="Prof Two", email="prof2@test.com",
        hashed_password="hash", role=UserRole.instructor,
    )
    student = User(
        id=3, name="Stu Test", email="stu@test.com",
        hashed_password="hash", role=UserRole.student,
    )
    lms_session = LMSSession(id=10, title="Week 1 Day 1")
    db.add_all([instructor, instructor2, student, lms_session])
    db.commit()

    instructor_token = create_access_token(1, UserRole.instructor)
    instructor2_token = create_access_token(2, UserRole.instructor)
    student_token = create_access_token(3, UserRole.student)

    with TestClient(app) as c:
        yield c, instructor_token, instructor2_token, student_token

    app.dependency_overrides.clear()


# ===========================================================================
# Upload
# ===========================================================================

class TestUploadLecture:

    def test_successful_upload_extracts_and_creates_chunks(self, client, db, tmp_path):
        """Real (unmocked) storage — extraction genuinely runs against the
        file actually saved to disk, not a mocked-away path."""
        c, instr_token, _, _ = client
        data = _make_pptx_bytes(with_notes=True)

        with patch("app.services.storage._storage_root", return_value=tmp_path):
            res = _post(c, instr_token, 10, "week1.pptx", data)

        assert res.status_code == 201
        body = res.json()
        assert body["original_filename"] == "week1.pptx"
        assert body["session_id"] == 10
        assert body["extracted"] is True
        assert body["extraction_error"] is None

        lecture = db.query(LectureFile).filter(LectureFile.session_id == 10).one()
        assert lecture.instructor_id == 1
        chunks = db.query(LectureChunk).filter(LectureChunk.lecture_file_id == lecture.id).all()
        assert len(chunks) >= 2  # at least one slide_text + one notes chunk
        sources = {c.source.value for c in chunks}
        assert "slide_text" in sources
        assert "notes" in sources

    def test_non_instructor_rejected(self, client):
        c, _, _, student_token = client
        data = _make_pptx_bytes()
        res = _post_mocked(c, student_token, 10, "week1.pptx", data)
        assert res.status_code == 403

    def test_unauthenticated_rejected(self, client):
        c, _, _, _ = client
        data = _make_pptx_bytes()
        res = c.post(
            "/api/v1/sessions/10/lectures",
            files={"file": ("week1.pptx", data, "application/octet-stream")},
        )
        assert res.status_code == 401

    def test_nonexistent_session_404(self, client):
        c, instr_token, _, _ = client
        data = _make_pptx_bytes()
        res = _post_mocked(c, instr_token, 999, "week1.pptx", data)
        assert res.status_code == 404

    def test_ppt_rejected_at_endpoint_with_clear_message(self, client):
        c, instr_token, _, _ = client
        res = _post_mocked(c, instr_token, 10, "week1.ppt", b"legacy binary garbage")
        assert res.status_code == 422
        assert "Legacy .ppt format is not supported" in res.json()["detail"]

    def test_non_pptx_extension_rejected(self, client):
        c, instr_token, _, _ = client
        res = _post_mocked(c, instr_token, 10, "week1.pdf", b"%PDF-1.4 fake")
        assert res.status_code == 422

    def test_duplicate_filename_in_session_rejected(self, client, db):
        c, instr_token, _, _ = client
        data = _make_pptx_bytes()
        first = _post_mocked(c, instr_token, 10, "week1.pptx", data)
        assert first.status_code == 201

        second = _post_mocked(c, instr_token, 10, "week1.pptx", data)
        assert second.status_code == 409

    def test_corrupted_pptx_still_creates_row_marked_failed(self, client, db, tmp_path):
        """Extraction failure must never silently drop the upload — the raw
        file and row are still created, marked failed with a real reason.
        Uses real (unmocked) storage so extraction genuinely fails on real
        garbage bytes on disk, not merely on a mocked-away file."""
        c, instr_token, _, _ = client
        garbage = b"not a real pptx file at all"

        with patch("app.services.storage._storage_root", return_value=tmp_path):
            res = _post(c, instr_token, 10, "corrupt.pptx", garbage)

        assert res.status_code == 201
        body = res.json()
        assert body["extracted"] is False
        assert body["extraction_error"] is not None

        lecture = db.query(LectureFile).filter(LectureFile.session_id == 10).one()
        assert lecture.extracted is False
        assert lecture.extraction_error is not None
        chunks = db.query(LectureChunk).filter(LectureChunk.lecture_file_id == lecture.id).all()
        assert chunks == []


# ===========================================================================
# List
# ===========================================================================

class TestListLectures:

    def test_instructor_can_list(self, client, db):
        c, instr_token, _, _ = client
        _post_mocked(c, instr_token, 10, "week1.pptx", _make_pptx_bytes())

        res = c.get(
            "/api/v1/sessions/10/lectures",
            headers={"Authorization": f"Bearer {instr_token}"},
        )
        assert res.status_code == 200
        assert len(res.json()) == 1

    def test_student_can_list_too(self, client, db):
        """Lecture files are supporting material, downloadable by students —
        list is NOT instructor-only."""
        c, instr_token, _, student_token = client
        _post_mocked(c, instr_token, 10, "week1.pptx", _make_pptx_bytes())

        res = c.get(
            "/api/v1/sessions/10/lectures",
            headers={"Authorization": f"Bearer {student_token}"},
        )
        assert res.status_code == 200
        assert len(res.json()) == 1

    def test_unauthenticated_rejected(self, client):
        c, _, _, _ = client
        res = c.get("/api/v1/sessions/10/lectures")
        assert res.status_code == 401


# ===========================================================================
# Download
# ===========================================================================

class TestDownloadLecture:

    def test_student_can_download(self, client, db, tmp_path):
        """Real (unmocked) storage round-trip: upload via the real
        save_lecture_file, then download and confirm byte-for-byte match."""
        c, instr_token, _, student_token = client
        data = _make_pptx_bytes()

        with patch("app.services.storage._storage_root", return_value=tmp_path):
            upload_res = _post(c, instr_token, 10, "week1.pptx", data)
            assert upload_res.status_code == 201
            lecture_id = upload_res.json()["id"]

            download_res = c.get(
                f"/api/v1/sessions/10/lectures/{lecture_id}/download",
                headers={"Authorization": f"Bearer {student_token}"},
            )
            assert download_res.status_code == 200
            assert download_res.content == data

    def test_nonexistent_lecture_404(self, client):
        c, instr_token, _, _ = client
        res = c.get(
            "/api/v1/sessions/10/lectures/999/download",
            headers={"Authorization": f"Bearer {instr_token}"},
        )
        assert res.status_code == 404

    def test_unauthenticated_rejected(self, client, db):
        c, instr_token, _, _ = client
        upload_res = _post_mocked(c, instr_token, 10, "week1.pptx", _make_pptx_bytes())
        lecture_id = upload_res.json()["id"]

        res = c.get(f"/api/v1/sessions/10/lectures/{lecture_id}/download")
        assert res.status_code == 401
