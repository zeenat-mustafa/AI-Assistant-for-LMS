"""
Tests for Phase 7, Sub-feature 7.6: practice quizzes.

In-memory DB, real router + real quiz_generator; notebook reading, Chroma
retrieval and the LLM call are faked, so no network and no quota. The real
Gemini verification runs live outside the default suite.

Every test that uses the DB also asserts the Grade table is byte-for-byte
unchanged (autouse _grades_untouched) — a quiz must never touch real grades.

Run with:
    cd backend
    python -m pytest tests/test_quiz.py -v
"""

import builtins
import io
import json
import random
import re
import tempfile
from pathlib import Path
from unittest.mock import MagicMock

import nbformat
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.models.grade import Grade
from app.models.lecture_chunk import ChunkSource, LectureChunk
from app.models.lecture_file import LectureFile
from app.models.quiz_attempt import QuizAttempt
from app.models.session import LMSSession
from app.models.submission import Submission
from app.models.submission_file import SubmissionFile
from app.models.unsolved_file import UnsolvedFile
from app.models.user import User, UserRole
from app.services import embeddings, llm_provider, quiz_generator
from app.services.auth import create_access_token
from app.services.lecture_extraction import LEGACY_PPT_ERROR
from app.services.llm_provider import LLMProviderError
from app.services.quiz_generator import (
    QUIZ_PROMPT_RULES,
    QuizScopeNotFoundError,
    build_quiz_prompt,
    parse_quiz_response,
    retrieve_chunks_for_quiz_scope,
    select_material,
    shuffle_question_options,
)

API = "/api/v1/quiz"
STUDENT_ID, INSTRUCTOR_ID, OTHER_STUDENT_ID = 1, 2, 3

LAB_TEXT_A = "A tool is simply a Python function decorated with @tool. " * 5
LAB_TEXT_B = "The agent needs to know which tools it can use before it can call them. " * 4
LAB_GAP = "@tool\ndef calculator(expression: str):\n    # TODO\n    return ______"
OTHER_NB_TEXT = "Pandas merge combines two DataFrames on a shared key column. " * 5
LECTURE_TEXT = "LangGraph stores information inside a shared State object passed between nodes. " * 4
OTHER_SESSION_LECTURE = "Gradient descent updates weights in the direction of the negative gradient. " * 4

NOTEBOOK_CELLS = {
    "lab3.ipynb": [
        {"type": "markdown", "content": LAB_TEXT_A, "heuristic_hint": False},
        {"type": "code", "content": "   ", "heuristic_hint": False},
        {"type": "code", "content": LAB_GAP, "heuristic_hint": True},
        {"type": "markdown", "content": LAB_TEXT_B, "heuristic_hint": False},
    ],
    "pandas.ipynb": [
        {"type": "markdown", "content": OTHER_NB_TEXT, "heuristic_hint": False},
    ],
}


# ===========================================================================
# Fixtures
# ===========================================================================

@pytest.fixture()
def db():
    import app.models  # noqa: F401

    engine = create_engine(
        "sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _fk(dbapi_conn, _record):
        cursor = dbapi_conn.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    session.add_all([
        User(id=STUDENT_ID, name="Stu", email="student@demo.com", hashed_password="h", role=UserRole.student),
        User(id=INSTRUCTOR_ID, name="Ins", email="instructor@demo.com", hashed_password="h",
             role=UserRole.instructor),
        User(id=OTHER_STUDENT_ID, name="Other", email="other@demo.com", hashed_password="h",
             role=UserRole.student),
        LMSSession(id=7, title="Week 4 Day 1"),
        LMSSession(id=5, title="Week 3 Day 1"),
    ])
    session.commit()
    session.add_all([
        UnsolvedFile(id=20, session_id=7, original_filename="Week 10_Lab3.ipynb", file_path="7/assignments/lab3.ipynb"),
        UnsolvedFile(id=12, session_id=5, original_filename="pandas_merge.ipynb", file_path="5/assignments/pandas.ipynb"),
        LectureFile(id=1, session_id=7, original_filename="week 10_day1.pptx", file_path="7/lectures/w10.pptx",
                    extracted=True),
        LectureFile(id=2, session_id=5, original_filename="week3.pptx", file_path="5/lectures/w3.pptx",
                    extracted=True),
    ])
    session.commit()
    session.add_all([
        LectureChunk(id=1, lecture_file_id=1, slide_number=4, source=ChunkSource.slide_text, chunk_index=0,
                     chunk_text=LECTURE_TEXT),
        LectureChunk(id=2, lecture_file_id=2, slide_number=2, source=ChunkSource.notes, chunk_index=0,
                     chunk_text=OTHER_SESSION_LECTURE),
        # Real grade data, so "unchanged" is never trivially 0 == 0.
        Submission(id=1, session_id=7, student_id=STUDENT_ID, original_filename="sub.ipynb",
                   uploaded_file_path="7/submissions/1/sub.ipynb"),
    ])
    session.commit()
    session.add(SubmissionFile(id=1, submission_id=1, matched_unsolved_file_id=20, original_filename="sub.ipynb",
                               extracted_ipynb_path="7/submissions/1/extracted/sub.ipynb", graded=True))
    session.commit()
    session.add(Grade(id=1, submission_file_id=1, score=7.5, feedback_text="Good work.",
                      rationale_json='{"criteria": []}'))
    session.commit()
    yield session
    session.close()
    Base.metadata.drop_all(engine)


def _grade_snapshot(session):
    session.expire_all()
    return [
        (g.id, g.submission_file_id, g.score, g.feedback_text, g.rationale_json, g.summary, g.graded_at)
        for g in session.query(Grade).order_by(Grade.id).all()
    ]


@pytest.fixture(autouse=True)
def _grades_untouched(db):
    before = _grade_snapshot(db)
    assert len(before) == 1
    yield
    assert _grade_snapshot(db) == before


@pytest.fixture()
def client(db):
    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _headers(user_id, role=UserRole.student):
    return {"Authorization": f"Bearer {create_access_token(user_id, role)}"}


@pytest.fixture()
def student_headers():
    return _headers(STUDENT_ID)


def _valid_quiz_json(prompt: str) -> str:
    ids = re.findall(r"^--- (S\d+) \|", prompt, re.MULTILINE)
    return json.dumps({"questions": [
        {
            "question": f"Question {i}?",
            "options": [f"right {i}", f"wrong a {i}", f"wrong b {i}", f"wrong c {i}"],
            "correct_option_index": 0,
            "source_id": ids[i % len(ids)],
        }
        for i in range(5)
    ]})


@pytest.fixture()
def fakes(monkeypatch):
    """Fake notebook reads (by filename), Chroma retrieval, and the LLM call."""
    state = {"prompts": [], "retrieve_calls": [], "notebook_paths": [], "strip_flags": [], "hits": [],
             "llm": _valid_quiz_json}

    def fake_structure(path, strip_images=False):
        state["notebook_paths"].append(path)
        state["strip_flags"].append(strip_images)
        cells = NOTEBOOK_CELLS.get(Path(path).name)
        if cells is None:
            return {"valid": False, "cells": [], "error": "File not found"}
        return {"valid": True, "cells": cells, "error": None}

    def fake_retrieve(query, session_id=None, top_k=5, min_similarity=0.35):
        state["retrieve_calls"].append({"query": query, "session_id": session_id, "top_k": top_k})
        return state["hits"]

    def fake_call_llm(prompt, purpose="fast"):
        state["prompts"].append((prompt, purpose))
        result = state["llm"]
        if isinstance(result, Exception):
            raise result
        return result(prompt)

    monkeypatch.setattr(quiz_generator, "extract_notebook_structure", fake_structure)
    monkeypatch.setattr(embeddings, "retrieve", fake_retrieve)
    monkeypatch.setattr(llm_provider, "call_llm", fake_call_llm)
    return state


def _pptx_bytes() -> bytes:
    prs = Presentation()
    for title, body in [
        ("ReAct agents", "An agent alternates between reasoning steps and tool calls until it can answer. " * 4),
        ("Tool registries", "A registry maps each tool name string to the Python function that runs it. " * 4),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _ipynb_bytes() -> bytes:
    nb = nbformat.v4.new_notebook()
    nb.cells = [
        nbformat.v4.new_markdown_cell("Dictionaries map keys to values and look them up in constant time. " * 5),
        nbformat.v4.new_code_cell("# TODO: complete\nresult = ______"),
        nbformat.v4.new_markdown_cell("A list comprehension builds a new list from an iterable in one line. " * 4),
    ]
    return nbformat.writes(nb).encode("utf-8")


def _generate(client, headers, body):
    return client.post(f"{API}/generate", json=body, headers=headers)


# ===========================================================================
# parse_quiz_response
# ===========================================================================

LOOKUP = {f"S{i}": {"chunk_text": "x", "citation": f"Week 4 Day 1 · file.ipynb · cell {i}"} for i in range(1, 4)}


def _q(**overrides):
    q = {"question": "What does @tool do?", "options": ["a", "b", "c", "d"], "correct_option_index": 1,
         "source_id": "S2"}
    q.update(overrides)
    return q


def _raw(questions):
    return json.dumps({"questions": questions})


class TestParseQuizResponse:

    def test_well_formed_response_parses_with_server_side_citation(self):
        result = parse_quiz_response(_raw([_q() for _ in range(5)]), LOOKUP)
        assert result["valid"] is True, result["error"]
        assert len(result["questions"]) == 5
        assert result["questions"][0] == {
            "question": "What does @tool do?", "options": ["a", "b", "c", "d"],
            "correct_option_index": 1, "source_citation": "Week 4 Day 1 · file.ipynb · cell 2",
        }
        assert "source_id" not in result["questions"][0]

    def test_code_fences_and_surrounding_text_are_stripped(self):
        raw = "Here is your quiz:\n```json\n" + _raw([_q() for _ in range(5)]) + "\n```\nEnjoy!"
        assert parse_quiz_response(raw, LOOKUP)["valid"] is True

    @pytest.mark.parametrize("count", [0, 4, 6])
    def test_rejects_wrong_question_count(self, count):
        result = parse_quiz_response(_raw([_q() for _ in range(count)]), LOOKUP)
        assert result["valid"] is False
        assert "exactly 5" in result["error"]

    @pytest.mark.parametrize("options", [["a", "b", "c"], ["a", "b", "c", "d", "e"], "abcd", None])
    def test_rejects_wrong_option_count(self, options):
        questions = [_q() for _ in range(4)] + [_q(options=options)]
        result = parse_quiz_response(_raw(questions), LOOKUP)
        assert result["valid"] is False
        assert "exactly 4 options" in result["error"]

    @pytest.mark.parametrize("options", [["a", "  ", "c", "d"], ["a", "B", "b ", "d"], ["a", 2, "c", "d"]])
    def test_rejects_blank_duplicate_or_non_text_options(self, options):
        result = parse_quiz_response(_raw([_q(options=options)] + [_q() for _ in range(4)]), LOOKUP)
        assert result["valid"] is False

    @pytest.mark.parametrize("bad_index", [True, False, 4, -1, "1", 1.0, None])
    def test_rejects_invalid_correct_option_index(self, bad_index):
        result = parse_quiz_response(_raw([_q(correct_option_index=bad_index)] + [_q() for _ in range(4)]), LOOKUP)
        assert result["valid"] is False
        assert "correct_option_index" in result["error"]

    def test_rejects_missing_correct_option_index(self):
        q = _q()
        del q["correct_option_index"]
        result = parse_quiz_response(_raw([q] + [_q() for _ in range(4)]), LOOKUP)
        assert result["valid"] is False
        assert "correct_option_index" in result["error"]

    @pytest.mark.parametrize("source_id", ["S99", "s2", "S2 ", 2, None])
    def test_rejects_unknown_or_malformed_source_id(self, source_id):
        result = parse_quiz_response(_raw([_q(source_id=source_id)] + [_q() for _ in range(4)]), LOOKUP)
        assert result["valid"] is False
        assert "source_id" in result["error"]

    def test_rejects_missing_source_id(self):
        q = _q()
        del q["source_id"]
        assert parse_quiz_response(_raw([q] + [_q() for _ in range(4)]), LOOKUP)["valid"] is False

    @pytest.mark.parametrize("raw", ["", "not json at all", json.dumps([_q() for _ in range(5)]), '{"quiz": []}'])
    def test_rejects_non_json_or_wrong_root(self, raw):
        assert parse_quiz_response(raw, LOOKUP)["valid"] is False


# ===========================================================================
# Prompt, material selection, shuffling
# ===========================================================================

class TestPromptAndMaterial:

    def test_build_quiz_prompt_labels_blocks_and_returns_matching_lookup(self):
        chunks = [
            {"chunk_text": "first text", "citation": "Week 4 Day 1 · lab.ipynb · cell 1 (markdown)"},
            {"chunk_text": "second text", "citation": "week.pptx · slide 4 (speaker notes)"},
        ]
        prompt, lookup = build_quiz_prompt(chunks, 'The topic "tools".')
        assert QUIZ_PROMPT_RULES in prompt
        assert 'Quiz scope: The topic "tools".' in prompt
        assert "--- S1 | Week 4 Day 1 · lab.ipynb · cell 1 (markdown) ---\nfirst text" in prompt
        assert "--- S2 | week.pptx · slide 4 (speaker notes) ---\nsecond text" in prompt
        assert set(lookup) == set(re.findall(r"^--- (S\d+) \|", prompt, re.MULTILINE)) == {"S1", "S2"}
        assert lookup["S2"] is chunks[1]

    def test_select_material_keeps_everything_under_budget_in_order(self):
        chunks = [{"chunk_text": str(i) * 10, "citation": str(i)} for i in range(6)]
        assert select_material(chunks, random.Random(1), budget=1000) == chunks

    def test_select_material_respects_budget_keeps_order_and_varies_by_attempt(self):
        chunks = [{"chunk_text": "x" * 100, "citation": str(i)} for i in range(50)]
        selections = []
        for seed in range(5):
            picked = select_material(chunks, random.Random(seed), budget=1000)
            assert sum(len(c["chunk_text"]) for c in picked) <= 1000
            assert len(picked) == 10
            positions = [int(c["citation"]) for c in picked]
            assert positions == sorted(positions)
            selections.append(tuple(positions))
        assert len(set(selections)) > 1

    def test_select_material_skips_a_chunk_larger_than_the_budget(self):
        chunks = [{"chunk_text": "x" * 5000, "citation": "big"}, {"chunk_text": "y" * 10, "citation": "small"}]
        assert [c["citation"] for c in select_material(chunks, random.Random(0), budget=100)] == ["small"]

    def test_shuffle_remaps_correct_option_index_to_the_same_option_text(self):
        questions = [{"question": f"q{i}", "options": ["right", "w1", "w2", "w3"], "correct_option_index": 0,
                      "source_citation": "c"} for i in range(5)]
        positions = set()
        for seed in range(20):
            for q in shuffle_question_options(questions, random.Random(seed)):
                assert sorted(q["options"]) == sorted(["right", "w1", "w2", "w3"])
                assert q["options"][q["correct_option_index"]] == "right"
                positions.add(q["correct_option_index"])
        assert positions == {0, 1, 2, 3}
        assert questions[0]["options"] == ["right", "w1", "w2", "w3"]  # input not mutated


# ===========================================================================
# retrieve_chunks_for_quiz_scope — each scope mode
# ===========================================================================

class TestRetrieveChunksPerScope:

    def test_assignment_file_reads_only_that_file_and_drops_hinted_gap_cells(self, db, fakes):
        chunks = retrieve_chunks_for_quiz_scope("assignment_file", {"unsolved_file_id": 20}, db)
        assert [c["chunk_text"] for c in chunks] == [LAB_TEXT_A, LAB_TEXT_B]
        assert [c["citation"] for c in chunks] == [
            "Week 4 Day 1 · Week 10_Lab3.ipynb · cell 1 (markdown)",
            "Week 4 Day 1 · Week 10_Lab3.ipynb · cell 4 (markdown)",  # 1-indexed; gap cell 3 dropped
        ]
        assert [Path(p).name for p in fakes["notebook_paths"]] == ["lab3.ipynb"]
        assert fakes["strip_flags"] == [True]
        assert fakes["retrieve_calls"] == []

    def test_session_reads_that_sessions_lectures_and_notebooks_only(self, db, fakes):
        chunks = retrieve_chunks_for_quiz_scope("session", {"session_id": 7}, db)
        assert [c["chunk_text"] for c in chunks] == [LECTURE_TEXT, LAB_TEXT_A, LAB_TEXT_B]
        assert chunks[0]["citation"] == "Week 4 Day 1 · week 10_day1.pptx · slide 4 (slide text)"
        assert fakes["retrieve_calls"] == []

    def test_multiple_sessions_is_the_union_in_requested_order(self, db, fakes):
        chunks = retrieve_chunks_for_quiz_scope("multiple_sessions", {"session_ids": [5, 7]}, db)
        assert [c["chunk_text"] for c in chunks] == [
            OTHER_SESSION_LECTURE, OTHER_NB_TEXT, LECTURE_TEXT, LAB_TEXT_A, LAB_TEXT_B,
        ]
        assert chunks[0]["citation"] == "Week 3 Day 1 · week3.pptx · slide 2 (speaker notes)"
        assert fakes["retrieve_calls"] == []

    def test_topic_searches_all_sessions_top_10_and_drops_hinted_notebook_hits(self, db, fakes):
        fakes["hits"] = [
            {"source_type": "lecture", "source_file_id": 1, "session_id": 7, "slide_number": 4,
             "source": "slide_text", "chunk_text": LECTURE_TEXT, "similarity": 0.7},
            {"source_type": "notebook", "source_file_id": 20, "session_id": 7, "cell_index": 2,
             "cell_type": "code", "chunk_text": LAB_GAP, "similarity": 0.6},
            {"source_type": "notebook", "source_file_id": 20, "session_id": 7, "cell_index": 3,
             "cell_type": "markdown", "chunk_text": LAB_TEXT_B, "similarity": 0.5},
            {"source_type": "lecture", "source_file_id": 999, "session_id": 7, "slide_number": 1,
             "source": "notes", "chunk_text": "stale vector", "similarity": 0.5},
        ]
        chunks = retrieve_chunks_for_quiz_scope("topic", {"topic_text": "tools"}, db)
        assert fakes["retrieve_calls"] == [{"query": "tools", "session_id": None, "top_k": 10}]
        assert [c["chunk_text"] for c in chunks] == [LECTURE_TEXT, LAB_TEXT_B]
        assert chunks[1]["citation"] == "Week 4 Day 1 · Week 10_Lab3.ipynb · cell 4 (markdown)"

    def test_uploaded_pptx_and_ipynb_are_extracted_from_bytes(self, db, fakes):
        pptx = retrieve_chunks_for_quiz_scope(
            "uploaded_file", {"original_filename": "mine.pptx", "file_type": "pptx"}, db, uploaded_bytes=_pptx_bytes(),
        )
        assert [c["citation"] for c in pptx] == ["mine.pptx · slide 1 (slide text)", "mine.pptx · slide 2 (slide text)"]
        assert "reasoning steps and tool calls" in pptx[0]["chunk_text"]

        # The real in-memory notebook extractor (not the fake), gap cell dropped.
        fakes["notebook_paths"].clear()
        nb = retrieve_chunks_for_quiz_scope(
            "uploaded_file", {"original_filename": "mine.ipynb", "file_type": "ipynb"}, db, uploaded_bytes=_ipynb_bytes(),
        )
        assert [c["citation"] for c in nb] == ["mine.ipynb · cell 1 (markdown)", "mine.ipynb · cell 3 (markdown)"]
        assert fakes["notebook_paths"] == [] and fakes["retrieve_calls"] == []

    @pytest.mark.parametrize("scope_type,detail", [
        ("assignment_file", {"unsolved_file_id": 404}),
        ("session", {"session_id": 404}),
        ("multiple_sessions", {"session_ids": [7, 404]}),
    ])
    def test_missing_ids_raise_not_found(self, db, fakes, scope_type, detail):
        with pytest.raises(QuizScopeNotFoundError):
            retrieve_chunks_for_quiz_scope(scope_type, detail, db)


# ===========================================================================
# Hint-cell filter + per-filter exclusion counts + image stripping in material
# ===========================================================================

# Verbatim patterns from the real course notebooks (surveyed in 7.6 review).
REAL_HINT_CELLS = [
    # Week 10_Lab3.ipynb cell 6 — the 12x "💡 Hint" form
    "## Create the Gemini LLM\n\nImport the required class and initialize the model.\n\n💡 Hint\n\n"
    "The class name starts with\n\nChatGoogle...",
    # Week 10_Lab3.ipynb cell 42 — "Hint:" alone on a line
    "## Challenge\n\nModify your memory so it only stores the latest **5** messages.\n\nHint:\n\n"
    "Python list slicing may be useful.",
    # Week10_Day2.ipynb cell 43 — inline parenthetical, capitalised
    "(Hint: compare `'CUST-123'` here against the sample data format from Section 5.)",
    # week10_day1_lab.ipynb cell 28 — inline parenthetical, lowercase, names the gap answer
    "- Otherwise, evaluate it (hint: Python's `eval()` works fine for this trusted lab context) and",
    "- Only allow digits, `+ - * / ( ) .` and spaces — reject anything else (hint: `re.fullmatch` with",
]


class TestHintFilter:

    @pytest.mark.parametrize("text", REAL_HINT_CELLS)
    def test_every_real_hint_pattern_is_excluded(self, text):
        assert quiz_generator.is_hint_cell(text) is True

    @pytest.mark.parametrize("text", [
        # Object_detection_by_using_YOLOv8.ipynb cell 8 — the one real 💡 non-hint
        "<h3 align=\"left\"><font color='#f79a05'>💡 Inspiration:</font></h3>",
        "Thinking about the hinterland of a dataset, not its chintz, helps.",
        "Group by splits the data into smaller groups before aggregating.",
    ])
    def test_non_hint_text_is_kept(self, text):
        assert quiz_generator.is_hint_cell(text) is False

    @pytest.mark.parametrize("text", ["There is a subtle hint of seasonality here.", "HINTS for later"])
    def test_ordinary_prose_using_the_word_is_excluded_by_design(self, text):
        # Decided in 7.6 review: one whole-word rule, no heading/label logic.
        # No real notebook uses the word in ordinary prose today.
        assert quiz_generator.is_hint_cell(text) is True

    def test_both_filters_count_separately_and_heuristic_wins_ties(self):
        cells = [
            {"type": "code", "content": "# TODO (hint: use a loop)\nx = ____", "heuristic_hint": True},
            {"type": "markdown", "content": REAL_HINT_CELLS[0], "heuristic_hint": False},
            {"type": "code", "content": "# TODO\ny = ____", "heuristic_hint": True},
            {"type": "markdown", "content": "A tool is a decorated function.", "heuristic_hint": False},
        ]
        exclusions = quiz_generator.new_exclusion_counts()
        chunks = quiz_generator._notebook_chunks(cells, "lab.ipynb", None, exclusions)
        assert [c["chunk_text"] for c in chunks] == ["A tool is a decorated function."]
        assert exclusions == {"heuristic_hint": 2, "hint_pattern": 1}

    def test_json_scope_modes_apply_and_count_the_hint_filter(self, db, fakes, monkeypatch):
        hint_cell = {"type": "markdown", "content": REAL_HINT_CELLS[3], "heuristic_hint": False}
        monkeypatch.setitem(NOTEBOOK_CELLS, "lab3.ipynb", NOTEBOOK_CELLS["lab3.ipynb"] + [hint_cell])
        monkeypatch.setitem(NOTEBOOK_CELLS, "pandas.ipynb", NOTEBOOK_CELLS["pandas.ipynb"] + [hint_cell])

        for scope_type, detail, expected in [
            ("assignment_file", {"unsolved_file_id": 20}, {"heuristic_hint": 1, "hint_pattern": 1}),
            ("session", {"session_id": 7}, {"heuristic_hint": 1, "hint_pattern": 1}),
            ("multiple_sessions", {"session_ids": [7, 5]}, {"heuristic_hint": 1, "hint_pattern": 2}),
        ]:
            exclusions = quiz_generator.new_exclusion_counts()
            chunks = retrieve_chunks_for_quiz_scope(scope_type, detail, db, exclusions=exclusions)
            assert exclusions == expected, scope_type
            assert all("eval()" not in c["chunk_text"] for c in chunks)
        assert set(fakes["strip_flags"]) == {True}

    def test_topic_mode_filters_hint_cells_and_uses_source_text_not_chroma_text(self, db, fakes, monkeypatch):
        hint_cell = {"type": "markdown", "content": REAL_HINT_CELLS[0], "heuristic_hint": False}
        monkeypatch.setitem(NOTEBOOK_CELLS, "lab3.ipynb", NOTEBOOK_CELLS["lab3.ipynb"] + [hint_cell])
        fakes["hits"] = [
            {"source_type": "notebook", "source_file_id": 20, "session_id": 7, "cell_index": 0,
             "cell_type": "markdown", "chunk_text": "![x](data:image/png;base64,AAAA) stale Chroma text",
             "similarity": 0.6},
            {"source_type": "notebook", "source_file_id": 20, "session_id": 7, "cell_index": 2,
             "cell_type": "code", "chunk_text": LAB_GAP, "similarity": 0.6},
            {"source_type": "notebook", "source_file_id": 20, "session_id": 7, "cell_index": 4,
             "cell_type": "markdown", "chunk_text": REAL_HINT_CELLS[0], "similarity": 0.6},
        ]
        exclusions = quiz_generator.new_exclusion_counts()
        chunks = retrieve_chunks_for_quiz_scope("topic", {"topic_text": "tools"}, db, exclusions=exclusions)
        assert [c["chunk_text"] for c in chunks] == [LAB_TEXT_A]
        assert exclusions == {"heuristic_hint": 1, "hint_pattern": 1}
        assert fakes["strip_flags"] == [True]

    def test_uploaded_notebook_is_image_stripped_and_hint_filtered(self, db, fakes):
        nb = nbformat.v4.new_notebook()
        nb.cells = [
            nbformat.v4.new_markdown_cell("Split-Apply-Combine groups rows, then aggregates each group. " * 6),
            nbformat.v4.new_markdown_cell("![split.png](data:image/png;base64," + "iVBORw0KGgo" * 5000 + ")"),
            nbformat.v4.new_markdown_cell(REAL_HINT_CELLS[1]),
            nbformat.v4.new_code_cell("# TODO\nresult = ______"),
        ]
        exclusions = quiz_generator.new_exclusion_counts()
        chunks = retrieve_chunks_for_quiz_scope(
            "uploaded_file", {"original_filename": "u.ipynb", "file_type": "ipynb"}, db,
            uploaded_bytes=nbformat.writes(nb).encode("utf-8"), exclusions=exclusions,
        )
        assert [c["chunk_text"] for c in chunks][1] == "[image omitted]"
        assert len(chunks) == 2
        assert exclusions == {"heuristic_hint": 1, "hint_pattern": 1}

    def test_generation_logs_both_exclusion_counts(self, client, fakes, student_headers, caplog, monkeypatch):
        hint_cell = {"type": "markdown", "content": REAL_HINT_CELLS[2], "heuristic_hint": False}
        monkeypatch.setitem(NOTEBOOK_CELLS, "lab3.ipynb", NOTEBOOK_CELLS["lab3.ipynb"] + [hint_cell])
        with caplog.at_level("INFO", logger="app.services.quiz_generator"):
            resp = _generate(client, student_headers, {"scope_type": "assignment_file", "unsolved_file_id": 20})
        assert resp.status_code == 201
        assert "scope=assignment_file heuristic_hint=1 hint_pattern=1" in caplog.text
        assert "CUST-123" not in fakes["prompts"][0][0]


# ===========================================================================
# POST /quiz/generate and /quiz/generate/upload
# ===========================================================================

class TestGenerateEndpoints:

    @pytest.mark.parametrize("body,expected,excluded", [
        ({"scope_type": "assignment_file", "unsolved_file_id": 20}, [LAB_TEXT_A, LAB_TEXT_B], [LECTURE_TEXT, LAB_GAP]),
        ({"scope_type": "session", "session_id": 7}, [LECTURE_TEXT, LAB_TEXT_A], [OTHER_NB_TEXT, LAB_GAP]),
        ({"scope_type": "multiple_sessions", "session_ids": [7, 5]}, [LECTURE_TEXT, OTHER_NB_TEXT], [LAB_GAP]),
        ({"scope_type": "topic", "topic_text": "  shared   state \n objects "}, [LECTURE_TEXT], [LAB_TEXT_A]),
    ])
    def test_each_json_scope_reaches_the_llm_with_its_chunk_set(
        self, client, db, fakes, student_headers, body, expected, excluded,
    ):
        fakes["hits"] = [{"source_type": "lecture", "source_file_id": 1, "session_id": 7, "slide_number": 4,
                          "source": "slide_text", "chunk_text": LECTURE_TEXT * 2, "similarity": 0.8}]
        resp = _generate(client, student_headers, body)
        assert resp.status_code == 201, resp.text
        assert len(fakes["prompts"]) == 1
        prompt, purpose = fakes["prompts"][0]
        assert purpose == "quiz_generation"
        for text in expected:
            assert text in prompt
        for text in excluded:
            assert text not in prompt

        data = resp.json()
        assert data["scope_type"] == body["scope_type"]
        assert len(data["questions"]) == 5
        assert all(len(q["options"]) == 4 for q in data["questions"])
        assert data["submitted"] is False
        assert "does not affect your real grades" in data["notice"]

        row = db.get(QuizAttempt, data["id"])
        assert row.student_id == STUDENT_ID and row.score is None and row.student_answers_json is None
        assert row.submitted_at is None and row.max_score == 5
        detail = json.loads(row.scope_detail)
        if body["scope_type"] == "topic":
            assert detail == {"topic_text": "shared state objects"}
            assert fakes["retrieve_calls"][0]["query"] == "shared state objects"
        else:
            assert detail == {k: v for k, v in body.items() if k != "scope_type"}

    def test_generate_response_never_contains_correct_option_index(self, client, fakes, student_headers):
        resp = _generate(client, student_headers, {"scope_type": "assignment_file", "unsolved_file_id": 20})
        assert resp.status_code == 201
        assert "correct_option_index" not in resp.text

        view = client.get(f"{API}/{resp.json()['id']}", headers=student_headers)
        assert view.status_code == 200
        assert "correct_option_index" not in view.text

    def test_stored_citations_come_from_real_metadata(self, client, db, fakes, student_headers):
        resp = _generate(client, student_headers, {"scope_type": "assignment_file", "unsolved_file_id": 20})
        citations = {q["source_citation"] for q in resp.json()["questions"]}
        assert citations <= {
            "Week 4 Day 1 · Week 10_Lab3.ipynb · cell 1 (markdown)",
            "Week 4 Day 1 · Week 10_Lab3.ipynb · cell 4 (markdown)",
        }

    def test_upload_generates_without_disk_writes_or_any_chroma_call(
        self, client, db, fakes, student_headers, monkeypatch, tmp_path,
    ):
        chroma_mocks = {name: MagicMock(side_effect=AssertionError(f"{name} must not be called"))
                        for name in ("retrieve", "get_chroma_collection", "embed_text", "upsert_chunk")}
        for name, mock in chroma_mocks.items():
            monkeypatch.setattr(embeddings, name, mock)

        from app.services import storage
        monkeypatch.setattr(storage.settings, "storage_root", str(tmp_path))
        storage_before = sorted(tmp_path.rglob("*"))

        write_attempts = []
        real_open = builtins.open

        def guarded_open(file, mode="r", *args, **kwargs):
            if any(flag in mode for flag in "wax+"):
                write_attempts.append((file, mode))
                raise AssertionError(f"disk write attempted: {file!r} mode={mode!r}")
            return real_open(file, mode, *args, **kwargs)

        def no_temp_files(*args, **kwargs):
            write_attempts.append(("tempfile", args))
            raise AssertionError("temp file creation attempted")

        with monkeypatch.context() as guard:
            guard.setattr(builtins, "open", guarded_open)
            guard.setattr(io, "open", guarded_open)
            for name in ("mkstemp", "NamedTemporaryFile", "TemporaryFile", "mkdtemp"):
                guard.setattr(tempfile, name, no_temp_files)
            guard.setattr(Path, "write_bytes", lambda *a, **k: no_temp_files())
            guard.setattr(Path, "write_text", lambda *a, **k: no_temp_files())

            pptx_resp = client.post(
                f"{API}/generate/upload", headers=student_headers,
                files={"file": ("my_notes.pptx", _pptx_bytes(), "application/octet-stream")},
            )
            nb_resp = client.post(
                f"{API}/generate/upload", headers=student_headers,
                files={"file": ("practice.ipynb", _ipynb_bytes(), "application/octet-stream")},
            )

        assert pptx_resp.status_code == 201, pptx_resp.text
        assert nb_resp.status_code == 201, nb_resp.text
        assert write_attempts == []
        for mock in chroma_mocks.values():
            mock.assert_not_called()
        assert sorted(tmp_path.rglob("*")) == storage_before
        assert "correct_option_index" not in pptx_resp.text

        assert "reasoning steps and tool calls" in fakes["prompts"][0][0]
        assert "Dictionaries map keys to values" in fakes["prompts"][1][0]
        assert "# TODO: complete" not in fakes["prompts"][1][0]
        rows = {r.scope_type: json.loads(r.scope_detail) for r in db.query(QuizAttempt).all()}
        assert json.loads(db.get(QuizAttempt, pptx_resp.json()["id"]).scope_detail) == {
            "original_filename": "my_notes.pptx", "file_type": "pptx",
        }
        assert rows["uploaded_file"]["file_type"] in {"pptx", "ipynb"}
        assert all("path" not in key for detail in rows.values() for key in detail)

    @pytest.mark.parametrize("filename,expected", [
        ("old.ppt", LEGACY_PPT_ERROR), ("notes.pdf", "Unsupported file type '.pdf'"), ("noext", "(none)"),
    ])
    def test_upload_rejects_other_file_types_before_any_llm_call(
        self, client, fakes, student_headers, filename, expected,
    ):
        resp = client.post(f"{API}/generate/upload", headers=student_headers,
                           files={"file": (filename, b"whatever", "application/octet-stream")})
        assert resp.status_code == 422
        assert expected in resp.json()["detail"]
        assert fakes["prompts"] == []

    def test_corrupt_upload_is_a_clean_422(self, client, fakes, student_headers):
        resp = client.post(f"{API}/generate/upload", headers=student_headers,
                           files={"file": ("broken.pptx", b"not a zip", "application/octet-stream")})
        assert resp.status_code == 422
        assert "Couldn't read the uploaded file" in resp.json()["detail"]
        assert fakes["prompts"] == []

    def test_thin_material_is_rejected_with_no_llm_call(self, client, db, fakes, student_headers, monkeypatch):
        monkeypatch.setitem(NOTEBOOK_CELLS, "lab3.ipynb", [
            {"type": "markdown", "content": "Short intro.", "heuristic_hint": False},
            {"type": "code", "content": LAB_GAP * 20, "heuristic_hint": True},
        ])
        resp = _generate(client, student_headers, {"scope_type": "assignment_file", "unsolved_file_id": 20})
        assert resp.status_code == 422
        assert resp.json()["detail"] == quiz_generator.NOT_ENOUGH_MATERIAL_MESSAGE
        assert fakes["prompts"] == []
        assert db.query(QuizAttempt).count() == 0

    def test_topic_with_no_matches_is_rejected_with_no_llm_call(self, client, fakes, student_headers):
        fakes["hits"] = []
        resp = _generate(client, student_headers, {"scope_type": "topic", "topic_text": "quantum baking"})
        assert resp.status_code == 422
        assert resp.json()["detail"] == quiz_generator.NO_TOPIC_MATCH_MESSAGE
        assert fakes["prompts"] == []

    @pytest.mark.parametrize("body", [
        {"scope_type": "session", "session_id": 404},
        {"scope_type": "assignment_file", "unsolved_file_id": 404},
    ])
    def test_unknown_scope_ids_are_404(self, client, fakes, student_headers, body):
        assert _generate(client, student_headers, body).status_code == 404
        assert fakes["prompts"] == []

    @pytest.mark.parametrize("failure", [
        LLMProviderError("Both Gemini models failed: quota_metric org_id 429"),
        RuntimeError("API key not valid"),
    ])
    def test_llm_failure_is_a_clean_502_and_stores_nothing(self, client, db, fakes, student_headers, failure):
        fakes["llm"] = failure
        resp = _generate(client, student_headers, {"scope_type": "session", "session_id": 7})
        assert resp.status_code == 502
        assert resp.json()["detail"] == quiz_generator.QUIZ_GENERATION_UNAVAILABLE_MESSAGE
        assert "quota" not in resp.text and "API key" not in resp.text
        assert db.query(QuizAttempt).count() == 0

    def test_invalid_model_output_is_a_clean_502_and_stores_nothing(self, client, db, fakes, student_headers):
        fakes["llm"] = lambda prompt: _valid_quiz_json(prompt).replace('"S1"', '"S404"')
        resp = _generate(client, student_headers, {"scope_type": "session", "session_id": 7})
        assert resp.status_code == 502
        assert resp.json()["detail"] == quiz_generator.QUIZ_INVALID_RESPONSE_MESSAGE
        assert db.query(QuizAttempt).count() == 0

    @pytest.mark.parametrize("body", [
        {"scope_type": "session"},
        {"scope_type": "session", "session_id": 7, "topic_text": "extra"},
        {"scope_type": "multiple_sessions", "session_ids": [7]},
        {"scope_type": "multiple_sessions", "session_ids": [7, 7]},
        {"scope_type": "topic", "topic_text": "   "},
        {"scope_type": "topic", "topic_text": "x" * 201},
        {"scope_type": "uploaded_file"},
        {"scope_type": "everything"},
    ])
    def test_invalid_request_bodies_are_422(self, client, fakes, student_headers, body):
        assert _generate(client, student_headers, body).status_code == 422
        assert fakes["prompts"] == []

    def test_instructor_and_anonymous_cannot_generate(self, client, fakes):
        body = {"scope_type": "session", "session_id": 7}
        assert _generate(client, _headers(INSTRUCTOR_ID, UserRole.instructor), body).status_code == 403
        assert client.post(f"{API}/generate", json=body).status_code == 401
        assert fakes["prompts"] == []


# ===========================================================================
# POST /quiz/{id}/submit, GET /quiz/{id}, GET /quiz/history
# ===========================================================================

def _new_attempt(client, headers):
    resp = _generate(client, headers, {"scope_type": "assignment_file", "unsolved_file_id": 20})
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _correct_answers(db, attempt_id):
    db.expire_all()
    return [q["correct_option_index"] for q in json.loads(db.get(QuizAttempt, attempt_id).questions_json)]


class TestSubmitAndHistory:

    @pytest.mark.parametrize("pattern,expected_score", [
        ("all_correct", 5), ("all_wrong", 0), ("mixed", 3),
    ])
    def test_submit_scores_correctly(self, client, db, fakes, student_headers, pattern, expected_score):
        attempt_id = _new_attempt(client, student_headers)
        correct = _correct_answers(db, attempt_id)
        wrong = [(c + 1) % 4 for c in correct]
        answers = {"all_correct": correct, "all_wrong": wrong,
                   "mixed": correct[:3] + wrong[3:]}[pattern]

        resp = client.post(f"{API}/{attempt_id}/submit", json={"answers": answers}, headers=student_headers)
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["score"] == expected_score and data["max_score"] == 5
        assert data["not_a_real_grade"] is True
        assert data["score_label"] == f"{expected_score}/5 (practice quiz — does not affect your real grades)"
        assert "does not affect your real grades" in data["notice"]
        assert [q["correct_option_index"] for q in data["questions"]] == correct
        assert [q["student_answer_index"] for q in data["questions"]] == answers
        assert sum(q["is_correct"] for q in data["questions"]) == expected_score

        db.expire_all()
        row = db.get(QuizAttempt, attempt_id)
        assert row.score == expected_score
        assert json.loads(row.student_answers_json) == answers
        assert row.submitted_at is not None

    def test_second_submit_is_409_and_never_rescores(self, client, db, fakes, student_headers):
        attempt_id = _new_attempt(client, student_headers)
        correct = _correct_answers(db, attempt_id)
        first = client.post(f"{API}/{attempt_id}/submit", json={"answers": correct}, headers=student_headers)
        assert first.json()["score"] == 5
        db.expire_all()
        submitted_at = db.get(QuizAttempt, attempt_id).submitted_at

        wrong = [(c + 1) % 4 for c in correct]
        second = client.post(f"{API}/{attempt_id}/submit", json={"answers": wrong}, headers=student_headers)
        assert second.status_code == 409

        db.expire_all()
        row = db.get(QuizAttempt, attempt_id)
        assert row.score == 5
        assert json.loads(row.student_answers_json) == correct
        assert row.submitted_at == submitted_at

    def test_concurrent_submit_loser_cannot_rescore(self, db, fakes):
        """Service-level race: the row was submitted after the router's check."""
        attempt = QuizAttempt(student_id=STUDENT_ID, scope_type="session", scope_detail='{"session_id": 7}',
                              questions_json=json.dumps([{"correct_option_index": 0}] * 5), max_score=5)
        db.add(attempt)
        db.commit()
        quiz_generator.submit_quiz_attempt(db, attempt, [0, 0, 0, 0, 0])
        with pytest.raises(quiz_generator.QuizAlreadySubmittedError):
            quiz_generator.submit_quiz_attempt(db, attempt, [1, 1, 1, 1, 1])
        db.expire_all()
        assert db.get(QuizAttempt, attempt.id).score == 5

    def test_other_student_cannot_submit_or_view(self, client, db, fakes, student_headers):
        attempt_id = _new_attempt(client, student_headers)
        other = _headers(OTHER_STUDENT_ID)
        assert client.post(f"{API}/{attempt_id}/submit", json={"answers": [0] * 5}, headers=other).status_code == 403
        assert client.get(f"{API}/{attempt_id}", headers=other).status_code == 403
        assert client.get(f"{API}/history", headers=other).json()["attempts"] == []
        db.expire_all()
        assert db.get(QuizAttempt, attempt_id).submitted_at is None

    def test_unknown_attempt_is_404(self, client, fakes, student_headers):
        assert client.post(f"{API}/999/submit", json={"answers": [0] * 5}, headers=student_headers).status_code == 404
        assert client.get(f"{API}/999", headers=student_headers).status_code == 404

    @pytest.mark.parametrize("answers", [[0] * 4, [0] * 6, [0, 0, 0, 0, 4], [0, 0, 0, 0, -1],
                                         [True, 0, 0, 0, 0], ["1", 0, 0, 0, 0]])
    def test_invalid_answers_are_422_and_do_not_submit(self, client, db, fakes, student_headers, answers):
        attempt_id = _new_attempt(client, student_headers)
        resp = client.post(f"{API}/{attempt_id}/submit", json={"answers": answers}, headers=student_headers)
        assert resp.status_code == 422
        db.expire_all()
        assert db.get(QuizAttempt, attempt_id).submitted_at is None

    def test_view_after_submit_includes_answers(self, client, db, fakes, student_headers):
        attempt_id = _new_attempt(client, student_headers)
        client.post(f"{API}/{attempt_id}/submit", json={"answers": _correct_answers(db, attempt_id)},
                    headers=student_headers)
        data = client.get(f"{API}/{attempt_id}", headers=student_headers).json()
        assert data["score"] == 5 and data["not_a_real_grade"] is True
        assert "correct_option_index" in data["questions"][0]

    def test_history_is_own_submitted_attempts_most_recent_first(self, client, db, fakes, student_headers):
        first = _new_attempt(client, student_headers)
        second = _new_attempt(client, student_headers)
        unsubmitted = _new_attempt(client, student_headers)
        other_attempt = _new_attempt(client, _headers(OTHER_STUDENT_ID))
        client.post(f"{API}/{first}/submit", json={"answers": _correct_answers(db, first)}, headers=student_headers)
        client.post(f"{API}/{second}/submit", json={"answers": [(c + 1) % 4 for c in _correct_answers(db, second)]},
                    headers=student_headers)
        client.post(f"{API}/{other_attempt}/submit", json={"answers": [0] * 5}, headers=_headers(OTHER_STUDENT_ID))

        resp = client.get(f"{API}/history", headers=student_headers)
        assert resp.status_code == 200
        data = resp.json()
        assert [a["attempt_id"] for a in data["attempts"]] == [second, first]
        assert unsubmitted not in [a["attempt_id"] for a in data["attempts"]]
        assert [a["score"] for a in data["attempts"]] == [0, 5]
        assert data["not_a_real_grade"] is True
        assert all(a["not_a_real_grade"] is True for a in data["attempts"])
        assert all("does not affect your real grades" in a["score_label"] for a in data["attempts"])
        assert "correct_option_index" not in resp.text

    def test_instructor_cannot_use_history_or_submit(self, client, fakes):
        headers = _headers(INSTRUCTOR_ID, UserRole.instructor)
        assert client.get(f"{API}/history", headers=headers).status_code == 403
        assert client.post(f"{API}/1/submit", json={"answers": [0] * 5}, headers=headers).status_code == 403
